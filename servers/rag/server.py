import glob
import json
import os
from pathlib import Path
from typing import List

from langchain_chroma import Chroma
import openpyxl
import xlrd
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from mcp.server import MCPServer
from sentence_transformers import SentenceTransformer
from pypdf import PdfReader

WORKSPACE = Path("/workspace").resolve()
PERSIST_DIRECTORY = WORKSPACE / "rag_db"
MODEL_ID = "Qwen/Qwen3-Embedding-0.6B"
MODEL_REVISION = "97b0c614be4d77ee51c0cef4e5f07c00f9eb65b3"
MODEL_PATH = Path("/models/Qwen3-Embedding-0.6B")
MODEL_MARKER = MODEL_PATH / ".model-complete"
DB_MODEL_METADATA = PERSIST_DIRECTORY / "embedding-model.json"
COLLECTION_NAME = "documents"
mcp = MCPServer("rag_server")


class LocalEmbeddings(Embeddings):
    def __init__(self) -> None:
        self.model = SentenceTransformer(str(MODEL_PATH), device="cpu", local_files_only=True)

    def embed_documents(self, texts: List[str]) -> List[List[float]]:
        return self.model.encode(texts, batch_size=8, normalize_embeddings=True, show_progress_bar=False).tolist()

    def embed_query(self, text: str) -> List[float]:
        return self.model.encode(text, normalize_embeddings=True, show_progress_bar=False).tolist()


_embeddings: LocalEmbeddings | None = None


def get_embeddings() -> LocalEmbeddings:
    global _embeddings
    if not MODEL_MARKER.is_file():
        raise RuntimeError(f"Local embedding model is incomplete: {MODEL_PATH}")
    if _embeddings is None:
        _embeddings = LocalEmbeddings()
    return _embeddings


def model_metadata() -> dict[str, str]:
    return {"model_id": MODEL_ID, "revision": MODEL_REVISION}


def verify_database_model(vectorstore: Chroma) -> None:
    expected = model_metadata()
    PERSIST_DIRECTORY.mkdir(parents=True, exist_ok=True)
    if DB_MODEL_METADATA.is_file():
        actual = json.loads(DB_MODEL_METADATA.read_text(encoding="utf-8"))
        if actual != expected:
            raise RuntimeError("RAG database uses a different embedding model; clear and re-index it.")
        return
    if vectorstore._collection.count() > 0:
        raise RuntimeError("Existing RAG data has no model metadata; clear and re-index it for migration.")
    DB_MODEL_METADATA.write_text(json.dumps(expected, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def get_vectorstore() -> Chroma:
    store = Chroma(collection_name=COLLECTION_NAME, embedding_function=get_embeddings(), persist_directory=str(PERSIST_DIRECTORY))
    verify_database_model(store)
    return store


def resolve_workspace_files(pattern: str) -> list[Path]:
    if Path(pattern).is_absolute():
        raise ValueError(f"Absolute paths are not allowed: {pattern}")
    matches = []
    for raw_match in glob.glob(str(WORKSPACE / pattern)):
        resolved = Path(raw_match).resolve()
        try:
            resolved.relative_to(WORKSPACE)
        except ValueError as error:
            raise ValueError(f"Path escapes the workspace: {pattern}") from error
        if resolved.is_file():
            matches.append(resolved)
    return matches


def load_file(file_path: Path):
    extension, path = file_path.suffix.lower(), str(file_path)
    if extension == ".pdf":
        return [Document(page_content=page.extract_text() or "", metadata={"source": path, "page": index}) for index, page in enumerate(PdfReader(path).pages)]
    if extension == ".docx":
        text = "\n".join(paragraph.text for paragraph in DocxDocument(path).paragraphs)
        return [Document(page_content=text, metadata={"source": path})]
    if extension == ".doc":
        raise ValueError("Legacy .doc is not supported; save it as .docx first.")
    if extension == ".xlsx":
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        return [Document(page_content="\n".join("\t".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True)), metadata={"source": path, "sheet": sheet.title}) for sheet in workbook.worksheets]
    if extension == ".xls":
        workbook = xlrd.open_workbook(path)
        return [Document(page_content="\n".join("\t".join(str(value) for value in sheet.row_values(row)) for row in range(sheet.nrows)), metadata={"source": path, "sheet": sheet.name}) for sheet in workbook.sheets()]
    return [Document(page_content=file_path.read_text(encoding="utf-8", errors="replace"), metadata={"source": path})]


@mcp.tool()
def add_documents(file_paths: List[str]) -> str:
    """Add workspace-relative PDF, Word, Excel, Markdown, or text files to RAG."""
    try:
        documents, source_count = [], 0
        for pattern in file_paths:
            for resolved_path in resolve_workspace_files(pattern):
                documents.extend(load_file(resolved_path))
                source_count += 1
        if not documents:
            return "No matching documents were found in /workspace."
        splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200, separators=["\n\n", "\n", "。", "、", " ", ""])
        chunks = splitter.split_documents(documents)
        get_vectorstore().add_documents(documents=chunks)
        return f"Added {len(chunks)} chunks from {source_count} files."
    except Exception as error:
        return f"Error adding documents: {error}"


@mcp.tool()
def query_knowledge_base(query: str, n_results: int = 5) -> str:
    """Search indexed documents for passages related to a query."""
    try:
        if not 1 <= n_results <= 20:
            return "n_results must be between 1 and 20."
        results = get_vectorstore().similarity_search(query, k=n_results)
        if not results:
            return "No matching passages found."
        sections = []
        for index, document in enumerate(results, start=1):
            source = Path(document.metadata.get("source", "Unknown")).name
            page = document.metadata.get("page", "N/A")
            sections.append(f"Source {index}: {source} (Page: {page})\n{document.page_content.replace(chr(10), ' ')}")
        return "\n\n".join(sections)
    except Exception as error:
        return f"Error querying database: {error}"


@mcp.tool()
def rag_status() -> str:
    """Report model and database readiness without loading the embedding model."""
    metadata = json.loads(DB_MODEL_METADATA.read_text(encoding="utf-8")) if DB_MODEL_METADATA.is_file() else None
    return json.dumps({"model_ready": MODEL_MARKER.is_file(), "model_id": MODEL_ID, "model_revision": MODEL_REVISION, "database_exists": (PERSIST_DIRECTORY / "chroma.sqlite3").is_file(), "database_model": metadata}, ensure_ascii=False)


@mcp.tool()
def clear_database(confirm: bool = False) -> str:
    """Delete all indexed documents. Set confirm=true explicitly."""
    if not confirm:
        return "Database was not cleared. Call again with confirm=true."
    try:
        get_vectorstore().delete_collection()
        DB_MODEL_METADATA.unlink(missing_ok=True)
        return "RAG database cleared."
    except Exception as error:
        return f"Error clearing database: {error}"


if __name__ == "__main__":
    os.chdir(WORKSPACE)
    mcp.run()
